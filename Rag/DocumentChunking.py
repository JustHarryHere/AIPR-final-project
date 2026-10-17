from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

HEADING_RE = re.compile(r"^#{1,6}\s+(.*)$", re.MULTILINE)
TABLE_RE = re.compile(r"^\s*\|")
META_KEYS = ("course_id", "doc_type", "version", "date", "status", "owner")
SLIDE_RE = re.compile(r"^Slide (\d+)\b")
SUPPORTED_EXTENSIONS = (".md", ".txt", ".pptx")
PIPELINE_VERSION = 2

@dataclass
class Document:
    content: str
    doc_id: str = ""
    title: str = ""
    source: str = ""
    metadata: dict = field(default_factory=dict)
    created_at: datetime = field(default_factory=datetime.now)


@dataclass
class DocumentChunk:
    content: str
    chunk_id: str = ""
    doc_id: str = ""
    chunk_index: int = 0
    metadata: dict = field(default_factory=dict)
    embedding: Optional[list[float]] = None


def chunk_document(document: Document, chunk_size: int = 512, overlap: int = 64) -> list[DocumentChunk]:
    """Strategy "fixed": character window with overlap (exactly as in class)."""
    step = chunk_size - overlap
    chunks = []
    for chunk_index, start in enumerate(range(0, len(document.content), step)):
        content = document.content[start:start + chunk_size]
        if not content.strip():
            continue
        chunks.append(
            DocumentChunk(
                content=content,
                chunk_id=f"{document.doc_id}-chunk-{chunk_index}",
                doc_id=document.doc_id,
                chunk_index=chunk_index,
                metadata={"source": document.source, "title": document.title,
                          "start": start, "strategy": "fixed", **document.metadata},
            )
        )
    return chunks


def tag_sections(document: Document, chunks: list[DocumentChunk]) -> None:
    """Nearest heading at/before the chunk start (as in class rag_pipeline.py)."""
    headings = [(m.start(), m.group(1).strip()) for m in HEADING_RE.finditer(document.content)]
    for chunk in chunks:
        start = chunk.metadata.get("start", 0)
        section = "unknown"
        for pos, title in headings:
            if pos <= start:
                section = title
            else:
                break
        chunk.metadata["section"] = section


@dataclass
class DocumentStore:
    documents: dict = field(default_factory=dict)
    chunks: dict = field(default_factory=dict)

    def add_chunk(self, chunk: DocumentChunk) -> str:
        self.chunks[chunk.chunk_id] = chunk
        return chunk.chunk_id

    def add(self, document: Document) -> list[str]:
        self.documents[document.doc_id] = document
        return [self.add_chunk(c) for c in chunk_document(document)]

    def get_document(self, doc_id: str) -> Document:
        return self.documents[doc_id]

    def delete_document(self, doc_id: str) -> None:
        self.documents.pop(doc_id, None)
        self.chunks = {cid: c for cid, c in self.chunks.items() if c.doc_id != doc_id}

    def list_documents(self) -> list[Document]:
        return list(self.documents.values())

    def count(self) -> int:
        return len(self.chunks)


# ------------------------------------------------------------ heading strategy
def _blocks(text: str, lo: int, hi: int) -> list[tuple[int, int]]:
    """(start, end) of each paragraph/list/table in text[lo:hi].

    A block never contains a blank line, and a table (lines starting with '|')
    is always its own block - this is what keeps the weight table in one piece.
    """
    blocks: list[tuple[int, int]] = []
    cur_start = cur_end = None
    cur_table = False
    pos = lo
    for line in text[lo:hi].splitlines(keepends=True):
        line_start = pos
        pos += len(line)
        if not line.strip():
            if cur_start is not None:
                blocks.append((cur_start, cur_end))
                cur_start = None
            continue
        is_table = bool(TABLE_RE.match(line))
        if cur_start is not None and is_table != cur_table:
            blocks.append((cur_start, cur_end))
            cur_start = None
        if cur_start is None:
            cur_start, cur_table = line_start, is_table
        cur_end = line_start + len(line.rstrip("\r\n"))
    if cur_start is not None:
        blocks.append((cur_start, cur_end))
    return blocks


def _slide_meta(document: Document, section: str) -> dict:
    """{"slide": N} for converted decks, so citations can say "slide 7"."""
    if document.metadata.get("source_format") == "pptx":
        m = SLIDE_RE.match(section)
        if m:
            return {"slide": int(m.group(1))}
    return {}


def chunk_by_heading(document: Document, max_chars: int = 800) -> list[DocumentChunk]:
    """Strategy "heading": one section -> one or more chunks, blocks never split.

    - Blocks are packed into a chunk up to max_chars; one oversized block
      (e.g. a long table) stays whole rather than being cut.
    - Chunk text starts with "<title> > <heading>" so it stands on its own
      (the heading words also help retrieval).
    - metadata keeps start/section like the class chunks, so tag_sections-style
      evidence trails and the catalog's source_section work unchanged.
    """
    text = document.content
    heads = list(HEADING_RE.finditer(text))
    spans: list[tuple[str, int, int]] = []
    if not heads:
        spans.append((document.title, 0, len(text)))
    else:
        if heads[0].start() > 0:
            spans.append((document.title, 0, heads[0].start()))
        for i, m in enumerate(heads):
            hi = heads[i + 1].start() if i + 1 < len(heads) else len(text)
            spans.append((m.group(1).strip(), m.end(), hi))

    chunks: list[DocumentChunk] = []
    for section, lo, hi in spans:
        packed: list[tuple[int, int]] = []

        def emit() -> None:
            if not packed:
                return
            s0, e1 = packed[0][0], packed[-1][1]
            idx = len(chunks)
            prefix = document.title if section == document.title else f"{document.title} > {section}"
            chunks.append(DocumentChunk(
                content=f"{prefix}\n{text[s0:e1]}",
                chunk_id=f"{document.doc_id}-chunk-{idx}",
                doc_id=document.doc_id,
                chunk_index=idx,
                metadata={"source": document.source, "title": document.title,
                          "start": s0, "section": section, "strategy": "heading",
                          **document.metadata, **_slide_meta(document, section)},
            ))
            packed.clear()

        for b in _blocks(text, lo, hi):
            if packed and b[1] - packed[0][0] > max_chars:
                emit()
            packed.append(b)
        emit()
    return chunks


def chunk(document: Document, strategy: str = "heading", **kw) -> list[DocumentChunk]:
    if strategy == "heading":
        return chunk_by_heading(document, **kw)
    if strategy == "fixed":
        chunks = chunk_document(document, **kw)
        tag_sections(document, chunks)
        return chunks
    raise ValueError(f"strategy must be 'heading' or 'fixed', got {strategy!r}")

def load_metadata(path: str | Path) -> list[dict]:
    with open(path, encoding="utf-8") as f:
        return json.load(f)

def _clean(text: str) -> str:
    return " ".join(text.replace("\v", " ").split())

def _table_to_markdown(table) -> str:
    rows = [[_clean(c.text).replace("|", "\\|") for c in r.cells] for r in table.rows]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)

def _iter_shapes(shapes):
    """Shapes in reading order (top-to-bottom, left-to-right), groups flattened."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in sorted(shapes, key=lambda s: (s.top or 0, s.left or 0)):
        try:
            is_group = sh.shape_type == MSO_SHAPE_TYPE.GROUP
        except NotImplementedError:          # exotic shape types in python-pptx
            is_group = False
        if is_group:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def pptx_to_markdown(path: str | Path, notes_label: str = "Ghi chú") -> str:
    """Convert a .pptx into markdown, one `## Slide N: <title>` section per slide.

    - Title placeholder -> heading; body/subtitle paragraphs -> "- " bullets
      (indented by level); free text boxes -> plain lines.
    - Tables -> markdown tables (the heading chunker never splits them).
    - Speaker notes -> a "<notes_label>: ..." paragraph INSIDE the slide's
      section (not a heading), so notes stay in the same chunk and citation.
    - Slide numbers are the real deck numbers (empty slides keep their number
      but produce no chunk).
    NOT handled: text inside pictures/charts/SmartArt (needs OCR); hidden
    slides are included.
    """
    from pptx import Presentation          # lazy: only needed for .pptx corpora

    prs = Presentation(str(path))
    out: list[str] = []
    for n, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = _clean(title_shape.text_frame.text) if title_shape is not None else ""
        blocks: list[str] = []
        for sh in _iter_shapes(slide.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                md = _table_to_markdown(sh.table)
                if md:
                    blocks.append(md)
            elif sh.has_text_frame:
                lines = []
                for p in sh.text_frame.paragraphs:
                    text = _clean(p.text)
                    if not text:
                        continue
                    if sh.is_placeholder or p.level > 0:
                        lines.append("  " * p.level + "- " + text)
                    else:
                        lines.append(text)
                if lines:
                    blocks.append("\n".join(lines))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = _clean(slide.notes_slide.notes_text_frame.text)
            if notes:
                blocks.append(f"{notes_label}: {notes}")
        if not blocks and not title:
            continue                         # truly empty slide
        out.append(f"## Slide {n}: {title}" if title else f"## Slide {n}")
        out.extend(blocks)
    return "\n\n".join(out) + "\n"


def read_source(path: str | Path) -> str:
    """Return the text of a corpus file as markdown/plain text, by extension."""
    ext = Path(path).suffix.lower()
    if ext in (".md", ".txt"):
        return Path(path).read_text(encoding="utf-8")
    if ext == ".pptx":
        return pptx_to_markdown(path)
    raise ValueError(f"unsupported file type {ext!r} for {Path(path).name}; "
                     f"supported: {', '.join(SUPPORTED_EXTENSIONS)}")


def load_documents(docs_dir: str | Path, metadata_path: str | Path) -> list[Document]:
    docs = []
    for meta in load_metadata(metadata_path):
        path = Path(docs_dir) / meta["file"]
        docs.append(Document(
            content=read_source(path), doc_id=meta["doc_id"], title=meta["title"],
            source=meta["file"],
            metadata={**{k: meta.get(k) for k in META_KEYS},
                      "source_format": path.suffix.lower().lstrip(".")},
        ))
    return docs

def chunk_corpus(docs_dir, metadata_path, strategy: str = "heading", **kw) -> list[DocumentChunk]:
    return [c for d in load_documents(docs_dir, metadata_path) for c in chunk(d, strategy, **kw)]
